"""Build the lab report from recorded unittest runs. Requires python-docx."""

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Отчет_ЛР2_Веселков_Лев.docx"

# Each failed test is listed once with its observed and required behavior.
PRICE = [
    ("test_express_cost_exceeds_regular_cost", "2600 руб. вместо цены выше 5200 руб."),
    ("test_express_fragile_cost_exceeds_regular_cost", "3270 руб. вместо цены выше 6540 руб."),
]
DISTANCE = [
    ("test_501_kilometres_require_two_days", "04.09 вместо 05.09.2026"),
    ("test_999_kilometres_require_two_days", "04.09 вместо 05.09.2026"),
    ("test_1001_kilometres_require_three_days", "05.09 вместо 06.09.2026"),
    ("test_4999_kilometres_require_ten_days", "12.09 вместо 13.09.2026"),
]
EXPRESS_DAYS = [
    ("test_short_express_delivery_takes_at_least_one_day", "03.09 вместо 04.09.2026"),
    ("test_500_kilometre_express_delivery_takes_one_day", "03.09 вместо 04.09.2026"),
    ("test_three_day_express_route_rounds_up_to_two_days", "04.09 вместо 05.09.2026"),
    ("test_five_day_express_route_rounds_up_to_three_days", "05.09 вместо 06.09.2026"),
]
NAN_CASES = [
    ("test_nan_weight_is_rejected", "FAIL: вес NaN принят; результат (700, '2026-09-04')"),
    ("test_nan_distance_is_rejected_without_exception", "ERROR: ValueError в строке 48 при int(NaN)"),
]
TYPES = [
    ("test_string_weight_is_rejected_without_exception", "ERROR: TypeError, строка 12, вес '1'"),
    ("test_string_distance_is_rejected_without_exception", "ERROR: TypeError, строка 12, дистанция '100'"),
    ("test_none_weight_is_rejected_without_exception", "ERROR: TypeError, строка 12, вес None"),
    ("test_fractional_distance_is_rejected", "FAIL: 100.5 км приняты, стоимость 702 руб."),
    ("test_boolean_weight_is_rejected", "FAIL: True принят за вес 1 кг, стоимость 700 руб."),
    ("test_boolean_distance_is_rejected", "FAIL: True принят за 1 км, стоимость 205 руб."),
    ("test_string_express_flag_is_rejected", "FAIL: строка 'False' включает экспресс; 350 руб., 03.09.2026"),
]


def text(doc, value, style=None):
    return doc.add_paragraph(value, style)


def code(doc, value):
    for line in value.splitlines():
        paragraph = doc.add_paragraph(style="Code")
        paragraph.add_run(line)


def heading(doc, value, level=1):
    doc.add_heading(value, level=level)


def page(doc, title):
    doc.add_page_break()
    heading(doc, title)


def table(doc, headers, rows, widths, code_column=False):
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    result._tbl.tblPr.append(borders)
    for row_index, values in enumerate([headers, *rows]):
        row = result.rows[0] if row_index == 0 else result.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EEF5" if row_index == 0 else "FFFFFF")
            properties.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(value))
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = row_index == 0
            if row_index and code_column and col_index == 0:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            if not code_column and col_index > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def main():
    final = json.loads((ROOT / "results/summary.json").read_text(encoding="utf-8"))
    baseline = json.loads((ROOT / "results/baseline.json").read_text(encoding="utf-8"))
    for stats in (final, baseline):
        for filename, expected in stats["source_sha256_lf"].items():
            actual = hashlib.sha256((ROOT / filename).read_text(encoding="utf-8").encode("utf-8")).hexdigest()
            if actual != expected:
                raise ValueError(f"Source changed since testing: {filename}")
    triangle = final["suites"]["triangle"]
    delivery = final["suites"]["delivery"]
    initial = baseline["suites"]["triangle"]
    failures = {entry["test"].rsplit(".", 1)[1] for entry in delivery["failures"] + delivery["errors"]}
    documented = {name for name, _ in PRICE + DISTANCE + EXPRESS_DAYS + NAN_CASES + TYPES}
    if failures != documented:
        raise ValueError("Report must describe every failed delivery test exactly")
    if (triangle["total"], triangle["passed"], delivery["total"], delivery["passed"],
            len(delivery["failures"]), len(delivery["errors"]), initial["passed"]) != (39, 39, 44, 25, 15, 4, 33):
        raise ValueError("Unexpected totals: review the report text before rebuilding")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    # Some Word base templates put a blue paragraph border under Title.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(22)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    style = doc.styles.add_style("Code", 1)
    style.font.name = "Consolas"
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.0
    doc.core_properties.author = "Веселков Лев"
    doc.core_properties.title = "Лабораторная работа 2 Юнит тестирование"
    doc.core_properties.subject = "Треугольники и расчёт доставки"
    doc.core_properties.created = datetime.fromisoformat(final["tested_at"]).astimezone(timezone.utc)

    text(doc, "Лабораторная работа 2", "Title")
    heading(doc, "Юнит тестирование проектов и автоматизация проверок")
    text(doc, "Выполнил Веселков Лев\nВетка Lab_2\nДата проверки " + datetime.fromisoformat(final["tested_at"]).strftime("%d.%m.%Y"))
    text(doc, "Проверены метод определения вида треугольника из лабораторной работы 1 и учебный модуль расчёта доставки. "
         "Для собственного проекта исправлены ошибки на граничных значениях float. Все 39 его тестов проходят. "
         "В исходном модуле доставки 19 из 44 проверок не проходят; ниже приведены причины и предложения исправлений.")
    heading(doc, "А Количество тестов и результаты")
    rows = []
    for label, suite in (("Треугольники до исправлений", initial), ("Треугольники после исправлений", triangle), ("Исходная доставка", delivery)):
        rows.append((label, suite["total"], suite["passed"], len(suite["failures"]), len(suite["errors"])))
    table(doc, ("Модуль", "Всего", "Успех", "FAIL", "ERROR"), rows, (3.2, 0.95, 0.95, 0.95, 0.95))
    text(doc, "Написано 83 отдельных тестовых метода: 39 для треугольников и 44 для доставки. "
         "Контрольный прогон исходной версии треугольников не считается дополнительным набором тестов. "
         "В итоговом общем прогоне 64 успешные проверки, 15 FAIL и 4 ERROR. Пропущенных тестов нет.")
    heading(doc, "Б Непройденные тесты доставки")
    text(doc, "Точное количество — 19. Из них 15 проверок завершились несовпадением фактического и ожидаемого результата "
         "(FAIL), ещё 4 — необработанным исключением (ERROR). Они относятся к четырём группам аномалий. "
         "Десять проверок зависят от восстановленных правил экспресс-доставки и округления сроков, девять проверяют неверный ввод.")
    text(doc, f"Среда: Python {final['python']}, unittest, локальное виртуальное окружение venv. "
         "Файл Delivery.py сохранён в исходном виде; падения воспроизводятся обычным запуском unittest.")

    page(doc, "Условия проверки и исправления своего проекта")
    heading(doc, "Восстановленные правила доставки", 2)
    text(doc, "Из кода восстановлены диапазоны 0.1–50 кг и 1–5000 км включительно. Базовая цена равна 200 руб. "
         "плюс 5 руб. за километр. До 5 кг включительно весовой надбавки нет; свыше 5 и менее 20 кг применяется "
         "коэффициент 1.2, от 20 кг — 1.5. После него добавляются 300 руб. за хрупкую либо 1000 руб. за опасную упаковку. "
         "Обычная упаковка надбавки не имеет.")
    text(doc, "Дата отправки 03.09.2026 зафиксирована в строке 39 и не считается ошибкой. Возвращаются целочисленная "
         "стоимость и строка даты. Для неверного ввода описание функции задаёт результат (-1, '0000-00-00').")
    text(doc, "Для проверки сроков принято: каждые начатые 500 км требуют одного дня. Экспресс сокращает срок вдвое "
         "с округлением вверх, но не менее чем до одного дня. По смыслу дополнительной услуги его цена должна быть "
         "выше обычной. Эти три правила — допущения, восстановленные по контексту; отдельной тарифной спецификации нет. "
         "Поэтому расхождения D1–D3 оцениваются относительно этих правил. Предлагаемый коэффициент 2.0 необходимо согласовать; "
         "тесты не утверждают его точное значение.")
    text(doc, "Аннотации и описание неверного ввода приняты как контракт: вес — конечное число int/float, "
         "расстояние — int, признак экспресса — bool. Значения bool не считаются весом или расстоянием, "
         "хотя в Python этот тип наследует int.")
    heading(doc, "Исправления метода треугольников", 2)
    text(doc, "В исходной версии четыре теста обнаружили ошибочную классификацию. Абсолютный допуск 1e-12 "
         "делал стороны 3e-15, 4e-15 и 5e-15 равными; относительный допуск объединял близкие различные значения. "
         "Сравнение заменено на точное равенство уже преобразованных значений float.")
    text(doc, "Ещё два теста выявили потерю маленькой стороны при сложении с большой. Условие существования "
         "для отсортированных сторон заменено на smallest > largest - middle. Для поддержанных после этого "
         "экстремальных отношений сторон координаты строятся на самой длинной нормированной стороне, "
         "что исключает нулевое основание после нормирования. Высота вычисляется устойчивой формой формулы Герона.")
    text(doc, "Проверены типы фигур, три положения равных сторон, границы неравенства, нечисловой ввод, NaN, "
         "бесконечности, масштабирование, пропорции сторон и журналирование. После исправления: 39 из 39. "
         "На поле 100 × 100 очень близкие вершины могут совпадать после округления до пикселей; тип от этого не меняется.")

    page(doc, "В Локализация аномалий доставки")
    text(doc, "Номера строк приведены для исходного Delivery.py из коммита 133f24a. "
         "Названия в таблицах точно соответствуют методам класса DeliveryTests в tests/test_delivery.py.")
    heading(doc, "D1 Экспресс уменьшает стоимость", 2)
    text(doc, "Строка 36: total_cost *= 0.5. При включении экспресса цена уменьшается вдвое, "
         "что противоречит принятому правилу платного ускорения. Не проходят две проверки.")
    table(doc, ("Непройденный тест", "Фактически и ожидалось"), PRICE, (4.3, 2.7), True)
    text(doc, "Предлагаемое исправление строки 36: total_cost *= 2.0. "
         "Значение 2.0 — вариант тарифа, а не установленное заданием требование. "
         "Минимальное проверяемое условие — экспресс дороже обычной доставки.")
    heading(doc, "D2 Неполный участок пути не учитывается", 2)
    text(doc, "Строка 41: days_needed = max(1, distance // 500). Целочисленное деление "
         "отбрасывает остаток пути. По принятой модели 500 км в день маршрут 501 км требует два дня, "
         "а исходник отводит один. Не проходят четыре проверки.")
    table(doc, ("Непройденный тест", "Фактически и ожидалось"), DISTANCE, (4.3, 2.7), True)
    text(doc, "Предлагаемая замена строки 41:")
    code(doc, "days_needed = max(1, (distance + 499) // 500)")
    text(doc, "Расстояния 500 и 1000 км остаются контрольными точками без изменения ожидаемого срока. "
         "Проверки 501, 999, 1001 и 4999 км отличают округление вверх от усечения.")

    page(doc, "Срок экспресса и проверка специальных чисел")
    heading(doc, "D3 Нулевой и заниженный срок экспресса", 2)
    text(doc, "Строка 44: days_needed = days_needed // 2. Деление срока в один день даёт ноль, "
         "а нечётные сроки округляются вниз. Не проходят четыре проверки. Дата 03.09.2026 "
         "означает доставку в день отправки, хотя принята минимальная длительность один день.")
    table(doc, ("Непройденный тест", "Фактически и ожидалось"), EXPRESS_DAYS, (4.3, 2.7), True)
    text(doc, "Предлагаемая замена строки 44:")
    code(doc, "days_needed = max(1, (days_needed + 1) // 2)")
    text(doc, "Для сроков 1, 3 и 5 дней получаются соответственно 1, 2 и 3 дня. "
         "Этот расчёт применяется после исправленного расчёта обычного срока в строке 41.")
    heading(doc, "D4 Не проверяется корректность входных значений", 2)
    text(doc, "Основная локализация — строка 12, где до арифметики отсутствует проверка типов и конечности чисел. "
         "Сравнения NaN с границами ложны: вес NaN принимается, а расстояние NaN приводит к ValueError "
         "в строке 48. В обоих случаях ожидался результат (-1, '0000-00-00').")
    table(doc, ("Непройденный тест", "Фактический результат"), NAN_CASES, (4.3, 2.7), True)
    text(doc, "Исправление: проверять типы и диапазоны до расчёта цены, использовать math.isfinite для веса. "
         "Полный вариант защитного блока и остальные семь проверок этой группы приведены на следующей странице.")

    page(doc, "Неверные типы входных данных")
    text(doc, "Продолжение D4. Строки 12 и 35. Строки и None вызывают TypeError при сравнении; "
         "дробная дистанция и bool ошибочно участвуют в вычислениях. Строка 'False' истинна в условии if is_express. "
         "Для каждого из семи тестов ожидается (-1, '0000-00-00').")
    table(doc, ("Непройденный тест", "Фактический результат"), TYPES, (4.3, 2.7), True)
    heading(doc, "Предлагаемый блок проверки для D4", 2)
    text(doc, "Добавить import math и поместить блок перед текущей строкой 12. "
         "Проверка диапазонов до math.isfinite также не допускает преобразование слишком большого int во float.")
    code(doc, 'invalid = (-1, "0000-00-00")\n'
         'if type(weight) not in (int, float):\n'
         '    return invalid\n'
         'if type(distance) is not int or type(is_express) is not bool:\n'
         '    return invalid\n'
         'if not (0.1 <= weight <= 50.0 and 1 <= distance <= 5000):\n'
         '    return invalid\n'
         'if not math.isfinite(weight):\n'
         '    return invalid')
    text(doc, "Группа D4 включает девять непройденных тестов: пять FAIL и четыре ERROR. "
         "Вместе D1–D4 объясняют все 19 непройденных проверок доставки без повторного подсчёта.")

    page(doc, "Воспроизведение и материалы для сдачи")
    heading(doc, "Запуск в виртуальном окружении", 2)
    text(doc, "Открыть корневую папку в PyCharm и выбрать venv\\Scripts\\python.exe как интерпретатор. "
         "Для тестов достаточно стандартной библиотеки Python. Команды PowerShell из корня проекта:")
    code(doc, "python -m venv venv\n"
         ".\\venv\\Scripts\\python.exe -m unittest tests.test_triangle -v\n"
         ".\\venv\\Scripts\\python.exe -m unittest tests.test_delivery -v\n"
         ".\\venv\\Scripts\\python.exe -m unittest discover -v")
    text(doc, "Треугольники завершаются статусом OK. Исходная доставка и общий запуск завершаются "
         "FAILED (failures=15, errors=4), код возврата 1. Это воспроизводимые результаты "
         "исследования учебного модуля. Декораторы пропуска или ожидаемой ошибки не применяются.")
    heading(doc, "Протоколы и исходные версии", 2)
    text(doc, "results/triangle_baseline.txt содержит шесть падений до исправлений. "
         "results/triangle.txt содержит успешный повторный прогон. results/delivery.txt содержит "
         "имена всех 44 тестов, 15 несовпадений и четыре трассировки исключений. "
         "baseline.json и summary.json сохраняют статистику, время проверки и SHA256 исходных файлов "
         "с переводами строк, нормализованными до LF.")
    code(doc, ".\\venv\\Scripts\\python.exe scripts/run_checks.py --baseline\n"
         ".\\venv\\Scripts\\python.exe scripts/run_checks.py")
    text(doc, "Исходное решение треугольников хранится в evidence/triangle_lab1.py: это копия triangle.py "
         "из коммита 1d54bae ветки Lab_1. Исправления находятся в triangle.py. "
         "Модуль Delivery.py и README преподавателя сохранены из коммита 133f24a ветки Lab_2. "
         "Предложения исправлений доставки в этом отчёте не применены к исходному файлу.")
    heading(doc, "Вывод", 2)
    text(doc, "Обязательный минимум 20 тестов для своего проекта и 10 для доставки выполнен: создано 39 и 44 соответственно. "
         "Найденные ошибки своего проекта исправлены и проверены повторно. Для доставки восстановлены требования, "
         "указаны допущения и локализованы четыре группы аномалий. Зафиксированы все 19 непройденных проверок, "
         "включая исключения, и предложены изменения соответствующих строк.")
    heading(doc, "Источник задания", 2)
    text(doc, "https://github.com/leuff92/PTPM-VKI/tree/Lab_2\n"
         "Версия задания 133f24ade02c2df6ed2b040b0b27651c48da2673")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
